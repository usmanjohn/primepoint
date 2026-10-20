"""The portfolio as an .xlsx file: one sheet per section of the page, same numbers."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEAD_FILL = PatternFill("solid", fgColor="F2C230")
MONEY = "#,##0"


def _sheet(wb, title, headers, rows, money_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
    for r in rows:
        ws.append(list(r))
    for col in money_cols:                      # 1-based column numbers
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = MONEY
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(h) + 4)
    ws.freeze_panes = "A2"
    return ws


def write(p, stream):
    wb = Workbook()
    wb.remove(wb.active)

    _sheet(wb, "Сводка", ["Показатель", "Значение", "Комментарий"],
           [("Контрагент", p["name"], p["contragentid"]),
            ("Дата", p["pdate"].strftime("%d.%m.%Y"), p["client_type"])]
           + [(k["label"], round(k["value"]), k["note"]) for k in p["kpis"]], money_cols=[2])

    L = p["loans"]
    _sheet(wb, "Кредиты", ["Продукт", "Срочные", "Реструкт.", "Просрочка", "Суд", "Брутто",
                           "Резерв", "Нетто"],
           [(r["label"], r["current"], r["restructured"], r["overdue"], r["court"], r["gross"],
             r["reserve"], r["net"]) for r in L["rows"]]
           + [("ИТОГО", L["current"], L["restructured"], L["overdue"], L["court"], L["gross"],
               L["reserve"], L["net"])], money_cols=range(2, 9))

    _sheet(wb, "Сделки", ["Сделка", "РНК", "Валюта", "Выдано (экв.)", "Остаток", "Просрочка",
                          "Дн. просрочки", "Проценты", "Неисп. линия", "Дата закрытия",
                          "Дней осталось", "Счета"],
           [(d["dealid"], d["rnk_dealid"], d["currencyid"], d["given"], d["balance"],
             d["overdue"], d["overdue_days"], d["interest"], d["unused"], d["close_date"],
             d["days_left"], ", ".join(d["loan_accounts"])) for d in p["deals"]],
           money_cols=[4, 5, 6, 8, 9])

    _sheet(wb, "РНК", ["РНК", "Сделок", "Выдано (экв.)", "Остаток", "Просрочка", "Неисп. линия",
                       "Последнее закрытие", "Дней осталось"],
           [(r["rnk_dealid"], r["deals"], r["given"], r["balance"], r["overdue"], r["unused"],
             r["last_close"], r["days_left"]) for r in p["rnk"]], money_cols=[3, 4, 5, 6])

    _sheet(wb, "Депозиты", ["Вид", "Сумма", "Счетов", "Доля, %"],
           [(r["label"], r["amount"], r["count"], r["share"]) for r in p["deposits"]["rows"]],
           money_cols=[2])

    _sheet(wb, "Внебаланс", ["Вид", "Сумма", "Счетов"],
           [(r["label"], r["amount"], r["count"]) for r in p["offbalance"]["rows"]],
           money_cols=[2])

    _sheet(wb, "Обороты", ["Месяц", "Поступления", "Документов", "Списания", "Документов",
                           "Сальдо", "Просрочка ОД, дн.", "Просрочка %, дн."],
           [(m["label"], m["in"], m["in_docs"], m["out"], m["out_docs"], m["net"],
             m["od_days"], m["pr_days"]) for m in p["turnover"]["months"]],
           money_cols=[2, 4, 6])

    _sheet(wb, "Счета", ["Счёт", "Бал. счёт", "Валюта", "TOBO", "Раздел", "Что это", "Сальдо (как в базе)"],
           [(a["accountno"], a["bal_acc"], a["currency_code"], a["toboid"], a["section"],
             a["label"], a["saldo"]) for a in p["accounts"]], money_cols=[7])

    wb.save(stream)
