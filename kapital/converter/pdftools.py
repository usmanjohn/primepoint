"""
The PDF tools themselves: bytes in, bytes out. No Django and no disk: nothing is ever saved.

Each tool needs one package. If it is not installed, the tool is switched off on the page
(see AVAILABLE) and the other tools keep working:
    pypdf       merge, pages, split, rotate          pip install pypdf
    pdf2docx    PDF → Word                            pip install pdf2docx
    pdfplumber  PDF → Excel (+ openpyxl, already in the project)   pip install pdfplumber
"""
import io
import re
import zipfile

try:
    from pypdf import PdfReader, PdfWriter
    from pypdf.errors import PdfReadError
except ImportError:                      # tool switched off on the page
    PdfReader = PdfWriter = None
    PdfReadError = Exception

try:
    from pdf2docx import Converter
except ImportError:
    Converter = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

AVAILABLE = {
    "merge": PdfReader is not None,
    "pages": PdfReader is not None,
    "split": PdfReader is not None,
    "rotate": PdfReader is not None,
    "word": Converter is not None and PdfReader is not None,
    "excel": pdfplumber is not None and PdfReader is not None,
}
PACKAGE = {"merge": "pypdf", "pages": "pypdf", "split": "pypdf", "rotate": "pypdf",
           "word": "pdf2docx", "excel": "pdfplumber"}


class ToolError(Exception):
    """A problem the user can fix: shown on the page in Russian."""


# ── helpers ─────────────────────────────────────────────────────────────────────────

def _reader(data):
    try:
        reader = PdfReader(io.BytesIO(data))
    except (PdfReadError, ValueError, OSError):
        raise ToolError("Файл не похож на PDF или повреждён.")
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")          # many bank PDFs are "protected" with an empty password
        except Exception:
            ok = 0
        if not ok:
            raise ToolError("PDF защищён паролем — снимите защиту и попробуйте снова.")
    return reader


def _bytes(writer):
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def parse_pages(spec, total):
    """'1-3, 5, 8-' -> [0, 1, 2, 4, 7, 8, 9] (0-based, in the order written).

    a-b   pages a to b (b < a goes backwards: '5-3' = 5, 4, 3)
    a-    page a to the end;  -b  first page to b
    Empty spec = all pages.
    """
    spec = (spec or "").replace(" ", "").replace(";", ",")
    if not spec:
        return list(range(total))
    pages = []
    for part in filter(None, spec.split(",")):
        m = re.fullmatch(r"(\d*)-(\d*)|(\d+)", part)
        if not m or part == "-":
            raise ToolError(f"Не понимаю «{part}». Пишите так: 1-3, 5, 8-")
        if m.group(3):
            a = b = int(m.group(3))
        else:
            a = int(m.group(1) or 1)
            b = int(m.group(2) or total)
        for n in (a, b):
            if not 1 <= n <= total:
                raise ToolError(f"Страницы {n} нет: в файле {total} стр.")
        step = 1 if b >= a else -1
        pages.extend(range(a - 1, b - 1 + step, step))
    return pages


def natural_key(name):
    """'файл 10.pdf' after 'файл 2.pdf'."""
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", name)]


# ── the tools ───────────────────────────────────────────────────────────────────────

def merge(files):
    """files = [(name, bytes), ...] -> one PDF, in natural order of the file names."""
    if len(files) < 2:
        raise ToolError("Выберите хотя бы два PDF.")
    writer = PdfWriter()
    for name, data in sorted(files, key=lambda f: natural_key(f[0])):
        try:
            reader = _reader(data)
        except ToolError as e:
            raise ToolError(f"{name}: {e}")
        for page in reader.pages:
            writer.add_page(page)
    return _bytes(writer)


def pages(data, keep="", delete=""):
    """Keep pages in the order written (this also reorders), then drop the ones in delete."""
    reader = _reader(data)
    total = len(reader.pages)
    order = parse_pages(keep, total)
    drop = set(parse_pages(delete, total)) if delete.strip() else set()
    order = [i for i in order if i not in drop]
    if not order:
        raise ToolError("После удаления не осталось ни одной страницы.")
    writer = PdfWriter()
    for i in order:
        writer.add_page(reader.pages[i])
    return _bytes(writer)


def split(data, every, stem):
    """Every N pages -> a separate PDF, all of them in one ZIP."""
    reader = _reader(data)
    total = len(reader.pages)
    if every < 1:
        raise ToolError("Число страниц в части должно быть 1 или больше.")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for start in range(0, total, every):
            end = min(start + every, total)
            writer = PdfWriter()
            for i in range(start, end):
                writer.add_page(reader.pages[i])
            label = f"{start + 1:03d}" if end - start == 1 else f"{start + 1:03d}-{end:03d}"
            z.writestr(f"{stem}_{label}.pdf", _bytes(writer))
    return out.getvalue()


def rotate(data, spec, angle):
    """Rotate the listed pages (empty = all) clockwise by 90 / 180 / 270."""
    if angle not in (90, 180, 270):
        raise ToolError("Угол поворота: 90, 180 или 270.")
    reader = _reader(data)
    chosen = set(parse_pages(spec, len(reader.pages)))
    writer = PdfWriter()
    for i, page in enumerate(reader.pages):
        if i in chosen:
            page.rotate(angle)
        writer.add_page(page)
    return _bytes(writer)


def to_word(data, spec=""):
    """PDF -> .docx. Works on PDFs with real text; a scan comes out as pictures."""
    total = len(_reader(data).pages)
    chosen = parse_pages(spec, total)
    out = io.BytesIO()
    cv = Converter(stream=data)
    try:
        cv.convert(out, pages=chosen)
    finally:
        cv.close()
    return out.getvalue()


def to_excel(data, spec="", one_sheet=False):
    """Every table found in the PDF -> .xlsx: a sheet per table, or all on one sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    total = len(_reader(data).pages)
    chosen = parse_pages(spec, total)
    wb = Workbook()
    wb.remove(wb.active)
    found = 0
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for i in chosen:
            for k, table in enumerate(pdf.pages[i].extract_tables(), start=1):
                rows = [[_cell(c) for c in row] for row in table if any(c not in (None, "") for c in row)]
                if not rows:
                    continue
                found += 1
                if one_sheet:
                    ws = wb["Таблицы"] if "Таблицы" in wb.sheetnames else wb.create_sheet("Таблицы")
                    if ws.max_row > 1:
                        ws.append([])
                    ws.append([f"Страница {i + 1}, таблица {k}"])
                    ws.cell(ws.max_row, 1).font = Font(bold=True)
                else:
                    ws = wb.create_sheet(f"стр{i + 1}_т{k}"[:31])
                for row in rows:
                    ws.append(row)
    if not found:
        raise ToolError("Таблиц не найдено. Если это скан (картинка), его нельзя прочитать без распознавания текста.")
    for ws in wb.worksheets:
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _cell(v):
    """Text from a PDF table cell: numbers like '1 234 567,89' become real numbers.

    Codes stay text: an account number (20 digits) or '007' would lose digits as a number.
    """
    if v is None:
        return None
    s = " ".join(str(v).split())
    num = s.replace("\u00a0", "").replace(" ", "")
    if num[:1] in ("\u00ad", "\u2212", "\u2013"):    # PDFs write minus as soft hyphen, − or –
        num = "-" + num[1:]
    digits = num.lstrip("-").split(",")[0].split(".")[0]
    if len(digits) > 15 or (len(digits) > 1 and digits.startswith("0")):
        return s
    if re.fullmatch(r"-?\d+([.,]\d+)?", num):
        try:
            return float(num.replace(",", ".")) if re.search(r"[.,]", num) else int(num)
        except ValueError:
            pass
    return s
